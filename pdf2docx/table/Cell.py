'''Table Cell object.'''

from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ALIGN_VERTICAL
        
from ..common.Element import Element
from ..layout.Layout import Layout
from ..common import docx


def set_cell_width(cell, width=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    if width:
        tcW.set(qn('w:w'), str(int(width)))  # 单位: 1/20 pt
        tcW.set(qn('w:type'), 'dxa')        # 可调节宽度
    else:
        tcW.set(qn('w:type'), 'auto')       # 自动宽度


def set_cell_autowrap(cell, autowrap=True):
    """
    设置表格单元格是否自动换行
    :param cell: docx.table._Cell 对象
    :param autowrap: True = 自动换行 (默认), False = 不换行
    """
    tcPr = cell._tc.get_or_add_tcPr()
    noWrap = tcPr.find(qn('w:noWrap'))
    if autowrap:
        # 移除 noWrap，恢复自动换行
        if noWrap is not None:
            tcPr.remove(noWrap)
    else:
        # 添加或设置 noWrap
        if noWrap is None:
            noWrap = OxmlElement('w:noWrap')
            tcPr.append(noWrap)
        noWrap.set(qn('w:val'), 'true')

class Cell(Layout):
    '''Cell object.'''
    def __init__(self, raw:dict=None):
        raw = raw or {}
        super().__init__()
        self.restore(raw) # restore blocks and shapes

        # more cell properties
        self.bg_color     = raw.get('bg_color', None) # type: int
        self.border_color = raw.get('border_color', (0,0,0,0)) # type: tuple [int]
        self.border_width = raw.get('border_width', (0,0,0,0)) # type: tuple [float]
        self.merged_cells = raw.get('merged_cells', (1,1)) # type: tuple [int]

    
    @property
    def text(self):
        '''Text contained in this cell.'''
        if not self: return None
        # NOTE: sub-table may exists in
        # fixme: prev code did `if block.is_text_block`, but sometimes
        # there is no `is_text_block` member; would be good to ensure
        # this member is always present and avoid use of `hasattr()`.
        return '\n'.join([block.text if hasattr(block, 'text') else '<NEST TABLE>'
                                for block in self.blocks])


    @property
    def working_bbox(self):
        '''Inner bbox with border excluded.'''
        x0, y0, x1, y1 = self.bbox
        w_top, w_right, w_bottom, w_left = self.border_width
        bbox = (x0+w_left/2.0, y0+w_top/2.0, x1-w_right/2.0, y1-w_bottom/2.0)
        return Element().update_bbox(bbox).bbox # convert to fitz.Rect


    def store(self):
        if not bool(self): return None
        res = super().store()
        res.update({
            'bg_color': self.bg_color,
            'border_color': self.border_color,
            'border_width': self.border_width,
            'merged_cells': self.merged_cells,
            'is_default': self.is_default
        })
        return res


    def plot(self, page):
        '''Plot cell and its sub-layout.'''
        super().plot(page)
        self.blocks.plot(page)


    def make_docx(self, table, indexes, is_default:bool):
        '''Set cell style and assign contents.

        Args:
            table (Table): ``python-docx`` table instance.
            indexes (tuple): Row and column indexes, ``(i, j)``.
        '''
        # set cell style, e.g. border, shading, cell width
        self._set_style(table, indexes)

        # ignore merged cells
        if not bool(self):  return

        # merge cells
        n_row, n_col = self.merged_cells
        i, j = indexes
        docx_cell = table.cell(i, j)
        if n_row*n_col != 1 and ((i+n_row-1) * table._column_count + j+n_col-1) < len(table._cells): # check whether index is over length of cells
            _cell = table.cell(i+n_row-1, j+n_col-1)
            try:
                docx_cell.merge(_cell)
            except Exception as e:
                def show(c):
                    return f'[_tc.top={c._tc.top} _tc.bottom={c._tc.bottom}]'
                raise Exception(f'Failed to merge docx_cell={show(docx_cell)} _cell={show(_cell)}. {i=} {j=} {n_row=} {n_col=}') from e

        # ---------------------
        # cell width (cell height is set by row height)
        # ---------------------
        # experience: width of merged cells may change if not setting width for merged cells
        x0, y0, x1, y1 = self.bbox
        docx_cell.width = Pt(x1-x0)

        if is_default:
             # 默认格式：设置自动换行 + 垂直居中
            set_cell_autowrap(docx_cell, autowrap=True)
            docx_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # insert contents
        # NOTE: there exists an empty paragraph already in each cell, which should be deleted
        # first to avoid unexpected layout. `docx_cell._element.clear_content()` works here.
        # But, docx requires at least one paragraph in each cell, otherwise resulting in a
        # repair error.
        if self.blocks:
            docx_cell._element.clear_content()
            if is_default:
                # 默认格式：水平居中
                self.blocks.set_text_center()
            self.blocks.make_docx(docx_cell)


    def _set_style(self, table, indexes):
        '''Set ``python-docx`` cell style, e.g. border, shading, width, row height,
        based on cell block parsed from PDF.

        Args:
            table (Table): ``python-docx`` table object.
            indexes (tuple): ``(i, j)`` index of current cell in table.
        '''
        i, j = indexes
        docx_cell = table.cell(i, j)
        n_row, n_col = self.merged_cells

        # ---------------------
        # border style
        # ---------------------
        # NOTE: border width is specified in eighths of a point, with a minimum value of
        # two (1/4 of a point) and a maximum value of 96 (twelve points)
        keys = ('top', 'end', 'bottom', 'start')
        kwargs = {}
        for k, w, c in zip(keys, self.border_width, self.border_color):
            # skip if width=0 -> will not show in docx
            if not w: continue

            hex_c = f'#{hex(c)[2:].zfill(6)}'
            kwargs[k] = {
                'sz': 8*w, 'val': 'single', 'color': hex_c.upper()
            }

        # merged cells are assumed to have same borders with the main cell
        for m in range(i, i+n_row):
            for n in range(j, j+n_col):
                if len(table._cells) > m * table._column_count + n: # check whether index is over length of cells
                    docx.set_cell_border(table.cell(m, n), **kwargs)

        # ---------------------
        # cell bg-color
        # ---------------------
        if self.bg_color is not None:
            docx.set_cell_shading(docx_cell, self.bg_color)

        # ---------------------
        # clear cell margin
        # ---------------------
        # NOTE: the start position of a table is based on text in cell, rather than
        # left border of table. They're almost aligned if left-margin of cell is zero.
        docx.set_cell_margins(docx_cell, start=0, end=0, left=0, right=0)

        # set vertical direction if contained text blocks are in vertical direction
        if self.blocks.is_vertical_text:
            docx.set_vertical_cell_direction(docx_cell)
