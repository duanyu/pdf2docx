# -*- coding: utf-8 -*-

'''Collection of :py:class:`~pdf2docx.layout.Section` instances.
'''

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_ROW_HEIGHT
from ..common.Collection import BaseCollection
from ..common.docx import reset_paragraph_format
from ..common.docx import make_table_floating, set_cell_text_vertical, set_cell_font, set_cell_margins, set_cell_width
from .Section import Section
from ..common import constants


class Sections(BaseCollection):

    def restore(self, raws:list):
        """Restore sections from source dicts."""        
        self.reset()
        for raw in raws:
            section = Section().restore(raw)
            self.append(section)
        return self
    

    def parse(self, **settings):
        '''Parse layout under section level.'''
        for section in self: section.parse(**settings)
        return self


    def make_docx(self, doc):
        '''Create sections in docx.'''
        if not self:
            # 有可能有float images，make之
            if self.parent.float_images:
                p = doc.add_paragraph()
                for image in self.parent.float_images:
                    image.make_docx(p)
            return

        # mark paragraph index before creating current page
        n = len(doc.paragraphs)

        # create floating tables
        for table_dict in self.parent.float_tables:
            x0, y0, x1, y1 = table_dict['bbox']
            table_text = ''.join([span['text'] for span in table_dict['spans']])
            table_text_color = table_dict['spans'][0]['color']
            r = table_text_color & 0xFF
            g = (table_text_color >> 8) & 0xFF
            b = (table_text_color >> 16) & 0xFF
            table_text_font_size = table_dict['spans'][0]['size']
            table_docx = doc.add_table(rows=1, cols=1)
            make_table_floating(table_docx, x=x0 * 20, y=y0 * 20)
            table_docx.cell(0, 0).text = table_text
            set_cell_text_vertical(table_docx.cell(0, 0), direction="btLr")
            set_cell_font(table_docx.cell(0, 0), 'Times New Roman', table_text_font_size, RGBColor(r, g, b))
            set_cell_margins(table_docx.cell(0, 0), top=0, start=0, bottom=0, end=0)
            for row in table_docx.rows:
                row.height = Pt(y1-y0)
                row.height_rule = WD_ROW_HEIGHT.EXACTLY
            set_cell_width(table_docx.cell(0, 0), (x1-x0)*20)

        def create_dummy_paragraph_for_section(section):
            before_enter_num, after_enter_num = 0, 0
            if section.before_space >= 10:
                before_enter_num = int(section.before_space / 10) - 1
                section.before_space = section.before_space - 10 * before_enter_num

            if before_enter_num > 0:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.line_spacing = Pt(10)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run('\n' * before_enter_num)
                run.font.size = Pt(10)
                run.font.name = 'Times New Roman'
            else:
                p = doc.add_paragraph()
                line_height = min(section.before_space, 11)
                pf = reset_paragraph_format(p, line_spacing=Pt(line_height))
                pf.space_after = Pt(section.before_space-line_height)

        # ---------------------------------------------------
        # first section
        # ---------------------------------------------------
        # vertical position: add dummy paragraph only if before space is required
        section = self[0]
        # 计算sec之间的距离【为了多栏内各栏的布局用】，可以理解为after space list
        section_vertical_gaps = []
        page_margin_last_y = self.parent.margin[3]
        page_bbox = self.parent.bbox
        for sec_i, sec in enumerate(self):
            sec_bbox = sec.bbox
            if sec_i == len(self)-1:
                # 最后一个，用整个页面来计算
                section_vertical_gaps.append(max(page_bbox[3]-sec_bbox[3]-page_margin_last_y, 0.0))
            else:
                later_sec_bbox = self[sec_i+1].bbox
                section_vertical_gaps.append(max(later_sec_bbox[1] - sec_bbox[3], 0.0))

        # 如果first section的before space比较充裕，而last section的after space不充裕，那么让出一些before space出来，防止太频繁的换页
        # 如果有float images则不处理，防止错位
        if section_vertical_gaps[-1] < 30 and not self.parent.float_images:
            section.before_space = max(section.before_space-20, 0.0)
        elif section_vertical_gaps[-1] < 50 and not self.parent.float_images:
            section.before_space = max(section.before_space-10, 0.0)

        if section.before_space > constants.MINOR_DIST:
            create_dummy_paragraph_for_section(section)

        # create first section
        if section.num_cols==2: 
            doc.add_section(WD_SECTION.CONTINUOUS)
        section.make_docx(doc, section_vertical_gaps[0])

        # ---------------------------------------------------
        # more sections
        # ---------------------------------------------------
        for sec_i, section in enumerate(self[1:]):
            # 之前遗漏了这个逻辑
            if section.before_space > constants.MINOR_DIST:
                create_dummy_paragraph_for_section(section)

            # create new section symbol
            doc.add_section(WD_SECTION.CONTINUOUS)

            # set after space of last paragraph to define the vertical
            # position of current section
            # NOTE: the after space doesn't work if last paragraph is 
            # image only (without any text). In this case, set after
            # space for the section break.
            p = doc.paragraphs[-2] # -1 is the section break
            if not p.text.strip() and 'graphicData' in p._p.xml:
                p = doc.paragraphs[-1]
            pf = p.paragraph_format
            pf.space_after = Pt(section.before_space)
            
            # section content
            section.make_docx(doc, section_vertical_gaps[sec_i+1])

        # ---------------------------------------------------
        # create floating images
        # ---------------------------------------------------
        # lazy: assign all float images to first paragraph of current page
        for image in self.parent.float_images:
            image.make_docx(doc.paragraphs[n])


    def plot(self, page):
        '''Plot all section blocks for debug purpose.'''
        for section in self: 
            for column in section:
                column.plot(page, stroke=(1,1,0), width=1.5) # column bbox
                column.blocks.plot(page) # blocks